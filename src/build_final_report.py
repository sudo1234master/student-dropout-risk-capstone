"""Build the consolidated AIM capstone final report as a Word document."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "final" / "Student_Dropout_Risk_Capstone_Final_Report.docx"

NAVY = "17324D"
TEAL = "157F78"
RED = "D4475C"
GOLD = "D2A52A"
PALE_TEAL = "EAF4F3"
PALE_BLUE = "EEF3F7"
PALE_RED = "FAECEF"
MID_GREY = "66788A"
LIGHT_GREY = "D9E1E8"
WHITE = "FFFFFF"
BLACK = "111111"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        element = margins.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_left_border(cell, color: str, size: int = 24) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_cell_text(cell, text: str, bold=False, color=BLACK, size=8.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        set_cell_text(header.cells[index], text, bold=True, color=WHITE, size=8.5)
        set_cell_shading(header.cells[index], NAVY)
        if widths:
            header.cells[index].width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for column_index, value in enumerate(values):
            set_cell_text(row.cells[column_index], value, color=BLACK, size=8.2)
            if row_index % 2 == 1:
                set_cell_shading(row.cells[column_index], PALE_BLUE)
            if widths:
                row.cells[column_index].width = Inches(widths[column_index])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def set_column_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.30)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in [
        ("Title", 28, NAVY),
        ("Subtitle", 13, MID_GREY),
        ("Heading 1", 18, BLACK),
        ("Heading 2", 13, BLACK),
        ("Heading 3", 10.5, BLACK),
    ]:
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    if "Caption Compact" not in document.styles:
        caption = document.styles.add_style("Caption Compact", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "Arial"
        caption.font.size = Pt(8)
        caption.font.italic = True
        caption.font.color.rgb = RGBColor.from_string(MID_GREY)
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(7)

    if "Small Note" not in document.styles:
        note = document.styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
        note.font.name = "Arial"
        note.font.size = Pt(8)
        note.font.color.rgb = RGBColor.from_string(MID_GREY)
        note.paragraph_format.space_after = Pt(4)

    if "Metric" not in document.styles:
        metric = document.styles.add_style("Metric", WD_STYLE_TYPE.PARAGRAPH)
        metric.font.name = "Arial"
        metric.font.size = Pt(18)
        metric.font.bold = True
        metric.font.color.rgb = RGBColor.from_string(TEAL)
        metric.paragraph_format.space_after = Pt(0)

    for section in document.sections:
        header = section.header.paragraphs[0]
        header.text = "AIM Capstone  |  Student Dropout Risk"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Arial"
            run.font.size = Pt(7.5)
            run.font.color.rgb = RGBColor.from_string(MID_GREY)
        add_page_number(section.footer.paragraphs[0])


def add_kicker(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(TEAL)


def add_callout(document: Document, title: str, body: str, fill=PALE_TEAL, accent=TEAL) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.84)
    set_cell_shading(cell, fill)
    set_cell_left_border(cell, accent)
    set_cell_margins(cell, 120, 170, 120, 150)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(accent)
    paragraph = cell.add_paragraph(body)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(BLACK)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_picture(document: Document, relative_path: str, caption: str, width=6.55) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(ROOT / relative_path), width=Inches(width))
    caption_paragraph = document.add_paragraph(caption, style="Caption Compact")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_metric_strip(document: Document, metrics: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=len(metrics))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, (value, label) in enumerate(metrics):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, PALE_TEAL if index % 2 == 0 else PALE_BLUE)
        set_cell_margins(cell, 130, 120, 130, 120)
        cell.text = ""
        value_p = cell.paragraphs[0]
        value_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value_p.paragraph_format.space_after = Pt(2)
        run = value_p.add_run(value)
        run.font.name = "Arial"
        run.font.size = Pt(17)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(TEAL)
        label_p = cell.add_paragraph(label)
        label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_p.paragraph_format.space_after = Pt(0)
        for run in label_p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(7.5)
            run.font.color.rgb = RGBColor.from_string(MID_GREY)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def build_report() -> None:
    task3 = json.loads((ROOT / "reports/task3_analysis_summary.json").read_text())
    task4 = json.loads((ROOT / "reports/task4_model_summary.json").read_text())
    task5 = json.loads((ROOT / "reports/task5_bias_fairness_summary.json").read_text())
    heldout = pd.read_csv(ROOT / "reports/tables/task4/heldout_test_metrics.csv")
    threshold_row = heldout[heldout["prediction_rule"].str.contains("threshold")].iloc[0]
    capacity = pd.read_csv(ROOT / "reports/tables/task4/top_30pct_capacity_metrics.csv")
    capacity_row = capacity[capacity["partition"] == "Held-out test"].iloc[0]
    gaps = pd.read_csv(ROOT / "reports/tables/task5/fairness_gap_metrics.csv")

    document = Document()
    style_document(document)
    document.core_properties.title = "Student Dropout Risk Capstone Final Report"
    document.core_properties.subject = "AIM capstone project"
    document.core_properties.author = "AIM Capstone Project"
    document.core_properties.keywords = "student dropout, machine learning, fairness, explainability"

    # Title page
    document.add_paragraph().paragraph_format.space_after = Pt(48)
    add_kicker(document, "AIM Final Capstone")
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(10)
    title_run = title.add_run("Predicting Student Dropout Risk")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(5)
    subtitle_run = subtitle.add_run(
        "An end-to-end, reproducible machine learning project for early student support"
    )
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(13)
    subtitle_run.font.italic = True
    subtitle_run.font.color.rgb = RGBColor.from_string(MID_GREY)
    document.add_paragraph().paragraph_format.space_after = Pt(32)
    add_metric_strip(
        document,
        [
            ("4,424", "STUDENTS"),
            ("0.23", "ALERT THRESHOLD"),
            ("83.8%", "DROPOUT RECALL"),
            ("72.9%", "TOP-30% CAPTURE"),
        ],
    )
    document.add_paragraph().paragraph_format.space_after = Pt(42)
    add_callout(
        document,
        "Decision recommendation",
        "Use the model only in a controlled, non-punitive support pilot with human review, correction and appeal routes, calibration checks, and subgroup monitoring. Do not use it for automated adverse decisions.",
        fill=PALE_TEAL,
    )
    paragraph = document.add_paragraph("Prepared for AIM  |  September 2026", style="Small Note")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    # Executive summary
    document.add_heading("Executive summary", level=1)
    document.add_paragraph(
        "This capstone develops a reproducible early-warning model that ranks university students by the probability of dropping out after first-semester results become available. The intended decision is narrowly defined: help a student-support team prioritize optional outreach when support capacity is limited. The prediction is not a diagnosis, a causal explanation, or a basis for exclusion."
    )
    document.add_paragraph(
        "The project uses the public UCI Predict Students' Dropout and Academic Success dataset, containing 4,424 records from one Portuguese higher-education institution. The target has three outcomes: Graduate, Dropout, and Enrolled. Six second-semester academic fields were excluded from the end-of-first-semester prediction scope to reduce temporal leakage. Demographic fields were excluded from model inputs and retained for auditing."
    )
    document.add_paragraph(
        "A Random Forest achieved the highest dropout-focused cross-validated F2 score. Its dropout-alert threshold of 0.23 was selected from out-of-fold training predictions among thresholds that met the predefined recall and precision constraints. On the locked 885-record test set, the alert achieved 83.8% dropout recall and 62.6% dropout precision. Separately, the highest-risk 30% of students captured 207 of 284 held-out dropouts, or 72.9%, at 77.8% precision."
    )
    add_table(
        document,
        ["Success measure", "Predefined target", "Held-out result", "Status"],
        [
            ["Dropout recall", "At least 75%", f"{threshold_row['dropout_recall']:.1%}", "Met"],
            ["Dropout precision", "At least 50%", f"{threshold_row['dropout_precision']:.1%}", "Met"],
            ["Top-30% dropout capture", "At least 60%", f"{capacity_row['dropout_capture_rate']:.1%}", "Met"],
        ],
        widths=[2.15, 1.55, 1.55, 0.85],
    )
    add_callout(
        document,
        "Why the 30% review scenario is used",
        "Thirty percent is an illustrative workload scenario rather than a verified staffing limit. It is close to the source dataset's 32.1% dropout prevalence, provides a transparent comparison with random selection, and produces a list large enough to test ranking value. A real institution must replace it with documented outreach capacity before deployment.",
    )
    document.add_heading("Report map", level=2)
    add_bullets(
        document,
        [
            "Problem framing and measurable success criteria",
            "Dataset provenance, quality, scope, and governance",
            "Exploratory analysis and prediction-time feature engineering",
            "Cross-validated modeling, held-out evaluation, and capacity analysis",
            "Explainability, fairness screening, limitations, and operational controls",
            "Reproducible open-source repository structure and execution guide",
        ],
    )
    document.add_page_break()

    # Problem framing
    document.add_heading("1  Problem framing and decision context", level=1)
    document.add_heading("Business problem", level=2)
    document.add_paragraph(
        "Universities often learn about withdrawal after the opportunity for early support has narrowed. The proposed system gives student-support teams a consistent ranking after first-semester academic information becomes available. Staff can review high-risk records, verify context, and offer tutoring, counseling, financial guidance, or other support."
    )
    document.add_heading("Machine learning task", level=2)
    document.add_paragraph(
        "The analytical task is supervised multiclass classification. The model estimates probabilities for Dropout, Enrolled, and Graduate. The probability of Dropout serves as the operational risk score, while the three-class output preserves the original target definition."
    )
    add_table(
        document,
        ["Design element", "Project decision"],
        [
            ["Unit of analysis", "One student record"],
            ["Prediction point", "After first-semester results, before second-semester information"],
            ["Target", "Dropout, Enrolled, or Graduate"],
            ["Operational score", "Predicted probability of Dropout"],
            ["Intended user", "Authorized student-support team"],
            ["Allowed action", "Human review and optional supportive outreach"],
            ["Prohibited action", "Automated exclusion, discipline, grading, aid removal, or denial of service"],
        ],
        widths=[1.65, 4.65],
    )
    document.add_heading("Predefined success criteria", level=2)
    document.add_paragraph(
        "Dropout recall is prioritized because a missed dropout may lose the opportunity for support. Precision remains a constraint because unnecessary outreach can create stigma, frustration, and avoidable workload. The top-30% capture measure evaluates ranking quality under an explicit planning scenario, independently of the probability threshold."
    )
    add_callout(
        document,
        "Interpretation of the numeric targets",
        "A 75% recall target means finding at least three of every four actual dropouts. A 50% precision floor means at least one in every two alerts should correspond to an eventual dropout. A 60% top-30% capture target asks the ranking to find three in five dropouts while reviewing less than one third of students, twice the capture expected from random selection at the same capacity.",
        fill=PALE_BLUE,
        accent=NAVY,
    )
    document.add_page_break()

    # Data understanding
    document.add_heading("2  Dataset and data understanding", level=1)
    document.add_heading("Source and selection rationale", level=2)
    document.add_paragraph(
        "The UCI dataset was selected because it is public, well documented, licensed for reuse under CC BY 4.0, directly aligned with student dropout prediction, and large enough for a held-out evaluation. It includes enrollment context, academic preparation, socioeconomic indicators, macroeconomic variables, first-semester performance, and the three-class student outcome. The public DOI supports traceable attribution."
    )
    add_table(
        document,
        ["Dataset characteristic", "Value"],
        [
            ["Repository", "UCI Machine Learning Repository"],
            ["Dataset", "Predict Students' Dropout and Academic Success"],
            ["DOI", "10.24432/C5MC89"],
            ["License", "Creative Commons Attribution 4.0"],
            ["Rows", "4,424 student records"],
            ["Original inputs", "36 predictors"],
            ["Target", "Dropout, Enrolled, Graduate"],
            ["Machine-readable missing values", "None reported in the downloaded CSV"],
            ["Exact duplicate rows", "None detected"],
        ],
        widths=[2.25, 4.05],
    )
    add_picture(
        document,
        "reports/figures/task3/01_target_distribution.png",
        "Figure 1. Outcome distribution across the full dataset. Source: UCI dataset; reproduced by the Task 3 workflow.",
        width=6.35,
    )
    document.add_heading("Quality and scope decisions", level=2)
    add_bullets(
        document,
        [
            "The source spelling Nacionality was standardized to Nationality.",
            "Occupation code 99 was converted to an explicit Unknown or not stated category.",
            "Six second-semester academic variables were excluded because they would be unavailable at the stated prediction point.",
            "Continuous outliers were retained because they were plausible values and tree-based models are comparatively robust to them.",
            "The dataset was split once into 3,539 development records and 885 held-out test records using stratification and random seed 42.",
        ],
    )
    add_callout(
        document,
        "Data limitation",
        "The records come from one Portuguese institution. Performance, calibration, and fairness may change at another institution or in a later cohort. Local validation is required before student-facing use.",
        fill=PALE_RED,
        accent=RED,
    )
    # EDA and features
    document.add_heading("3  Exploratory analysis and feature engineering", level=1)
    document.add_paragraph(
        "Training-only exploratory analysis showed that first-semester academic progress and immediate financial indicators have the clearest descriptive relationship with dropout. These patterns guide feature construction and interpretation, but they do not establish causation."
    )
    add_picture(
        document,
        "reports/figures/task3/02_key_factor_dropout_rates.png",
        "Figure 2. Training-partition dropout rates for selected descriptive factors. Associations are not causal effects.",
        width=6.55,
    )
    add_table(
        document,
        ["Engineered feature", "Definition and purpose"],
        [
            ["First-semester approval rate", "Approved units divided by enrolled units, bounded from zero to one"],
            ["Evaluation intensity", "Evaluations divided by enrolled units"],
            ["No-evaluation rate", "Units without evaluation divided by enrolled units"],
            ["Progress index", "Approval rate multiplied by normalized first-semester grade"],
            ["Academic preparation gap", "Admission grade minus previous qualification grade"],
            ["No-enrollment flag", "Marks records with no first-semester enrolled units"],
            ["First-choice flag", "Marks first-choice application order"],
            ["Financial-risk flag", "Combines debtor status or tuition not current"],
            ["Nontraditional-age flag", "Marks age 25 or older for descriptive audit only"],
        ],
        widths=[2.2, 4.1],
    )
    document.add_heading("Feature screening and dimensionality", level=2)
    document.add_paragraph(
        "The primary prediction dataset contains 30 original first-semester predictors and nine engineered features. After audit-only demographic fields were removed, the model used 31 source columns. One-hot encoding produced 125 transformed features. Mutual information and a diagnostic Extra Trees model supported a 30-feature screening option, which Task 4 compared with using all encoded features inside cross-validation."
    )
    document.add_paragraph(
        "Principal component analysis required 61 components to explain 80% of variance, 75 for 90%, and 84 for 95%. The first two components explained only 9.4%, so PCA was retained as a diagnostic artifact rather than the default modeling path."
    )
    # Modeling
    document.add_heading("4  Model implementation and evaluation", level=1)
    document.add_heading("Experimental design", level=2)
    document.add_paragraph(
        "Task 4 compared a dummy prior baseline with Logistic Regression, Decision Tree, Random Forest, and Histogram Gradient Boosting. Each candidate used an end-to-end pipeline for imputation, categorical encoding, numeric scaling, optional feature selection, and classification. Five-fold stratified cross-validation and grid search used the development set only."
    )
    add_picture(
        document,
        "reports/figures/task4/01_cross_validated_model_comparison.png",
        "Figure 3. Five-fold cross-validated model comparison on the development partition.",
        width=6.55,
    )
    document.add_paragraph(
        "Random Forest produced the highest cross-validated dropout F2 and was selected according to the predefined dropout-focused objective. Logistic Regression showed a slightly stronger macro F1, which demonstrates that model choice depends on the decision objective rather than a single universal score."
    )
    document.add_heading("Threshold selection", level=2)
    document.add_paragraph(
        "The alert threshold was selected from out-of-fold training probabilities. Among thresholds satisfying at least 75% dropout recall and 50% precision, 0.23 maximized dropout F2. No held-out test information influenced this choice. The threshold changes the alert decision while retaining the same underlying three-class probabilities."
    )
    add_picture(
        document,
        "reports/figures/task4/02_threshold_tradeoff.png",
        "Figure 4. Training-only threshold trade-off using out-of-fold probabilities.",
        width=6.35,
    )
    document.add_heading("Held-out performance", level=2)
    add_metric_strip(
        document,
        [
            (f"{threshold_row['dropout_recall']:.1%}", "DROPOUT RECALL"),
            (f"{threshold_row['dropout_precision']:.1%}", "DROPOUT PRECISION"),
            (f"{threshold_row['dropout_f2']:.1%}", "DROPOUT F2"),
            (f"{threshold_row['accuracy']:.1%}", "OVERALL ACCURACY"),
        ],
    )
    document.add_paragraph(
        "The thresholded alert found 238 of 284 held-out dropouts and generated 142 false-positive alerts. This trade-off is consistent with the support objective: recall increases relative to default argmax prediction, while precision and overall multiclass accuracy decrease. False positives still matter, so outreach should remain low burden and supportive."
    )
    add_picture(
        document,
        "reports/figures/task4/05_heldout_capacity_and_calibration.png",
        "Figure 5. Held-out capacity ranking and probability calibration diagnostics.",
        width=6.55,
    )
    add_table(
        document,
        ["Capacity result", "Held-out value"],
        [
            ["Review list", f"{int(capacity_row['students_flagged'])} of 885 students"],
            ["Dropouts captured", f"{int(capacity_row['dropouts_captured'])} of {int(capacity_row['total_dropouts'])}"],
            ["Dropout capture rate", f"{capacity_row['dropout_capture_rate']:.1%}"],
            ["Precision at capacity", f"{capacity_row['precision_at_capacity']:.1%}"],
            ["Lift over random selection", f"{capacity_row['lift_at_capacity']:.2f} times"],
        ],
        widths=[3.1, 3.2],
    )
    document.add_heading("Interpretation", level=2)
    document.add_paragraph(
        "The model met all three predefined held-out criteria. The result supports continued evaluation, not unrestricted deployment. The test set is a single historical split from the same institution, and the probability calibration chart indicates additional calibration work before a risk score is presented as an absolute likelihood."
    )
    # Explainability and fairness
    document.add_heading("5  Explainability, ethical AI, and fairness", level=1)
    document.add_heading("Global and local explanations", level=2)
    document.add_paragraph(
        "Tree SHAP was calculated for all 885 held-out records. The largest grouped contributors were first-semester progress index, approval rate, tuition status, approved units, and the financial-risk flag. Local examples cover a true positive, false positive, and false negative. SHAP explains how the model produced a score; it does not show why a student leaves or what intervention will change the outcome."
    )
    add_picture(
        document,
        "reports/figures/task5/01_shap_global_importance.png",
        "Figure 6. Global Tree SHAP importance for the Dropout class on held-out records.",
        width=6.55,
    )
    document.add_paragraph(
        "Partial dependence and individual conditional expectation curves were generated for progress index and approval rate. Because engineered features are related to their component variables, the curves are diagnostic model-response views rather than feasible or causal intervention simulations."
    )
    document.add_heading("Fairness audit", level=2)
    document.add_paragraph(
        "Gender, age, and international status were excluded from the model inputs and used only for auditing. Debtor status, tuition status, and scholarship status were audited as incomplete socioeconomic proxies. Race is absent, and nationality or international status was not treated as a substitute for race."
    )
    add_picture(
        document,
        "reports/figures/task5/07_fairness_gap_summary.png",
        "Figure 7. Fairness screening gaps at the locked 0.23 dropout-alert threshold.",
        width=6.45,
    )
    flagged = gaps[gaps["equalized_odds_review_flag"]].copy()
    flag_rows = []
    for _, row in flagged.iterrows():
        flag_rows.append(
            [
                str(row["audit_dimension"]),
                str(row["group"]),
                f"{float(row['equalized_odds_max_gap']):.1%}",
                "Review",
            ]
        )
    add_table(
        document,
        ["Audit dimension", "Compared group", "Maximum TPR/FPR gap", "Screen"],
        flag_rows,
        widths=[1.65, 2.85, 1.25, 0.7],
    )
    document.add_paragraph(
        f"The audit produced {task5['fairness_results']['four_fifths_review_flags']} four-fifths screening flags and {task5['fairness_results']['equalized_odds_review_flags']} equalized-odds screening flags. The largest TPR/FPR gap was {task5['fairness_results']['largest_equalized_odds_gap']['gap']:.1%} for {task5['fairness_results']['largest_equalized_odds_gap']['group']} relative to {task5['fairness_results']['largest_equalized_odds_gap']['reference_group']}."
    )
    add_callout(
        document,
        "How to interpret the fairness results",
        "The flags identify comparisons that require review. They are not legal findings and do not prove discrimination. Historical outcome differences may reflect unequal access or institutional conditions. The practical benefits and burdens of outreach, false positives, and missed students must be evaluated with affected stakeholders.",
        fill=PALE_RED,
        accent=RED,
    )
    document.add_heading("Key ethical limitations", level=2)
    add_bullets(
        document,
        [
            "Race is unavailable, so racial fairness cannot be assessed.",
            "Gender is binary in the source data and may not represent student identity.",
            "The international group has only 24 held-out records and the scholarship group has only 19 actual dropouts, producing unstable estimates.",
            "Financial indicators do not fully measure income, wealth, caregiving, accessibility costs, or access to institutional support.",
            "Historical labels can encode structural and institutional inequities.",
            "Explanations and subgroup relationships are descriptive rather than causal.",
        ],
    )
    # Operations and governance
    document.add_heading("6  Operational recommendation and governance", level=1)
    document.add_paragraph(
        "The model is suitable only for a controlled support pilot after local review. The pilot should test whether outreach improves access to support and student outcomes without imposing disproportionate burden. A current institutional cohort should validate performance, calibration, and fairness before any scale-up."
    )
    add_table(
        document,
        ["Control", "Required practice", "Evidence to monitor"],
        [
            ["Support-only policy", "Limit alerts to offers of counseling, tutoring, and assistance", "Contact reason and service offered"],
            ["Human review", "Confirm context before student contact", "Overrides and review notes"],
            ["Explanation and appeal", "Provide plain-language reasons and a correction route", "Corrections, complaints, and appeal outcomes"],
            ["Calibration", "Recalibrate on current institutional data", "Brier score and calibration by cohort"],
            ["Fairness monitoring", "Track alert rate, recall, false-positive rate, and precision by group", "Confidence intervals and sample sizes"],
            ["Outcome monitoring", "Measure uptake, workload, retention, and unintended harm", "Pilot impact and burden indicators"],
            ["Stop conditions", "Pause when performance, fairness, or governance limits fail", "Documented thresholds and accountable owner"],
        ],
        widths=[1.35, 3.0, 2.0],
    )
    document.add_heading("Pilot sequence", level=2)
    add_bullets(
        document,
        [
            "Prepare: confirm governance ownership, current data definitions, capacity, and prohibited uses.",
            "Validate: evaluate the locked workflow on a later local cohort and recalibrate probabilities.",
            "Pilot: use one cohort with limited counselors and optional, low-burden outreach.",
            "Measure: compare support uptake, workload, retention, calibration, and subgroup outcomes.",
            "Gate: advance, revise, or stop using predefined performance and fairness thresholds.",
        ],
    )
    add_callout(
        document,
        "Financial value remains unquantified",
        "The public dataset does not include intervention cost or retention value. Return on investment should be estimated only after the institution supplies support costs and the pilot measures credible retention uplift.",
        fill=PALE_BLUE,
        accent=NAVY,
    )
    document.add_page_break()

    # Reproducibility and repository
    document.add_heading("7  Reproducibility and open-source repository", level=1)
    document.add_paragraph(
        "Task 7 consolidates the user-edited and locally verified deliverables into an open-source-style repository. The code and notebooks use project-root discovery, so they can run from the repository root or the notebooks directory. Exact tested package versions are recorded in requirements.txt."
    )
    add_table(
        document,
        ["Path", "Contents"],
        [
            ["src/", "Reusable Task 3-5 workflows, notebook builders, execution helpers, validators, and the pipeline runner"],
            ["notebooks/", "Executed analytical notebooks and the Task 6 technical presentation notebook"],
            ["data/", "Public raw data, UCI metadata, processed fixed splits, and the data dictionary workbook"],
            ["models/", "Saved preprocessing, PCA, candidate models, locked model bundle, and SHAP array"],
            ["configs/", "Saved modeling and fairness-audit configurations"],
            ["reports/", "Task reports, HTML exports, JSON summaries, figures, tables, and this final report"],
            ["presentations/", "Business PowerPoint and technical HTML presentation"],
        ],
        widths=[1.25, 5.1],
    )
    document.add_heading("Reproduction commands", level=2)
    code_lines = [
        "python3 -m venv .venv",
        "source .venv/bin/activate",
        "python -m pip install -r requirements.txt",
        "python src/run_pipeline.py --validate-only",
        "python src/run_pipeline.py",
        "python src/run_pipeline.py --execute-notebooks",
    ]
    table = document.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F4F6F8")
    set_cell_margins(cell, 130, 150, 130, 150)
    cell.text = ""
    for index, line in enumerate(code_lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(NAVY)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    document.add_heading("Reproducibility controls", level=2)
    add_bullets(
        document,
        [
            "Python 3.13.9 and exact tested library versions are documented.",
            "Random seed 42 and the fixed stratified split are saved.",
            "Preprocessing and feature selection are fitted inside model pipelines.",
            "Model comparison and threshold tuning use training data only.",
            "The locked test predictions, metrics, SHAP values, and fairness tables are committed.",
            "Validators confirm artifact completeness, calculations, notebook execution, and model-audit alignment.",
        ],
    )
    document.add_paragraph(
        "The repository includes an MIT license for original code, separate UCI dataset attribution under CC BY 4.0, a contribution guide, a citation note, and a gitignore file that excludes local environments, caches, and notebook checkpoints."
    )
    document.add_page_break()

    # Conclusion and references
    document.add_heading("8  Conclusion", level=1)
    document.add_paragraph(
        "The project demonstrates a complete machine learning workflow from problem framing and public-data selection through exploratory analysis, model comparison, held-out evaluation, explainability, fairness screening, stakeholder communication, and reproducible packaging. The selected Random Forest met the predefined held-out recall, precision, and ranking targets."
    )
    document.add_paragraph(
        "The technical result is promising for further study, but the governance conclusion is deliberately cautious. Fairness gaps, incomplete sensitive-attribute coverage, calibration needs, one-institution data, and uncertain intervention impact prevent unrestricted deployment. The next justified step is a monitored support pilot on a new local cohort, with explicit stop conditions and no automated adverse decisions."
    )
    add_callout(
        document,
        "Final decision",
        "Proceed only to local validation and a controlled support pilot. Treat the risk score as one input to human review, measure whether support helps, and stop or revise the system when performance, fairness, or governance conditions are not met.",
        fill=PALE_TEAL,
    )
    document.add_heading("References", level=1)
    references = [
        "Realinho, V., Machado, J., Baptista, L., and Martins, M. V. (2021). Predict Students' Dropout and Academic Success. UCI Machine Learning Repository. https://doi.org/10.24432/C5MC89",
        "UCI Machine Learning Repository. Dataset license: Creative Commons Attribution 4.0 International.",
        "Project artifacts. Task 1 Problem Understanding and Framing; Task 2 Data Collection and Understanding; Task 3 EDA and Feature Engineering; Task 4 Model Implementation; Task 5 Bias and Fairness Analysis; Task 6 Stakeholder Communication.",
        "scikit-learn documentation for cross-validation, pipeline construction, classification metrics, calibration, and model inspection.",
        "SHAP documentation and TreeExplainer implementation used for model-behavior explanations.",
    ]
    for item in references:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.16)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(item)

    document.add_heading("Artifact statement", level=2)
    document.add_paragraph(
        "This report synthesizes the project-specific analysis, tables, figures, code outputs, and author-edited task reports in this repository. Public dataset facts and third-party software are attributed above. All numerical results are traceable to committed JSON summaries and CSV tables."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
